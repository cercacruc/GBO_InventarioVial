"""Read-only extraction from the supplied SCAP workbook. No Excel dependency at runtime.
Run from repository root: python tools/extract_scap_catalog.py
Requires openpyxl for development only. Source cells and SHA-256 accompany every generated catalog.
"""
import argparse
import hashlib
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/reference/SCAP_PUENTE_AGUA_BLANCA.xlsx'
ASSETS = ROOT / 'app/src/main/assets'
LIST_RANGES = {
    'routeType':'R51:R53', 'over':'C49:C56', 'alignment':'G49:G51',
    'service':'J49:J53', 'environment':'N49:N51', 'spans':'C133:C134',
    'category':'C136:C139', 'edge':'C144:C150', 'predominant':'G144:G150',
    'slabMaterial':'J144:J148', 'wearingSurface':'N144:N148',
    'beamType':'E153:E156', 'beamMaterial':'H153:H157', 'beamShape':'L153:L156',
    'abutmentElevation':'C161:C166', 'abutmentMaterial':'G161:G164',
    'abutmentFoundation':'J161:J164', 'foundationMaterial':'N161:N165',
    'pierElevation':'C205:C208', 'pierMaterial':'G205:G209',
    'pierFoundation':'L205:L208', 'pierFoundationMaterial':'P205:P208',
    'anchorElevation':'C215:C217', 'anchorMaterial':'G215:G216',
    'anchorFoundation':'L215:L216', 'anchorFoundationMaterial':'P215:P216',
    'railingType':'C282:C285', 'railingMaterial':'G282:G285', 'railingSecondary':'U282:U287',
    'sidewalkMaterial':'J282:J284', 'bearingType':'N282:N287', 'bearingMaterial':'R282:R285',
    'jointType':'C290:C293', 'jointMaterial':'F290:F293', 'drainType':'I290:I293',
    'drainMaterial':'M290:M293', 'accessAlignment':'R290:R293',
    'visibility':'C296:C298', 'yesNo':'G296:G297', 'parallelPossibility':'N364:N366',
    'informative':'G360:G360', 'preventive':'G362:G364', 'regulatory':'G367:G369',
    'horizontal':'G372:G372', 'parallelSubstructure':'G375:G376', 'roadCondition':'J375:J378',
    'fordSoil':'N375:N380', 'foundationSoil':'H434:H440',
}
TYPE_RANGES = {'DEFINITIVO':'E134:E140','PROVISIONALES':'H134:H136',
               'ALCANTARILLA':'L134:L138','ARTESANALES':'P134:P138'}

def dump(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')

def extract(source=SOURCE):
    workbook=openpyxl.load_workbook(source, data_only=True)
    formulas=openpyxl.load_workbook(source, data_only=False)
    fingerprint=hashlib.sha256(source.read_bytes()).hexdigest()
    aux=workbook['AUXILIAR']; elements=[]
    groups={'SUPERESTRUCTURA','SUBESTRUCTURA','DETALLES','CAUCE','ACCESOS'}
    for row in range(5, aux.max_row+1):
        code=aux.cell(row,3).value
        if not isinstance(code,(int,float)): continue
        factor=aux.cell(row,6).value
        assert isinstance(factor,(int,float)) and 0<=factor<=1, (row,factor)
        item={'code':str(int(code)), 'name':str(aux.cell(row,4).value).strip(),
              'unit':str(aux.cell(row,5).value).strip(), 'importanceFactor':factor,
              'group':str(aux.cell(row,7).value).strip(), 'source':f'AUXILIAR!C{row}:G{row}'}
        assert item['group'] in groups, item
        elements.append(item)
    assert elements and len({e['code'] for e in elements})==len(elements)
    dump(ASSETS/'scap_elements.json',{'sourceSha256':fingerprint,'elements':elements})
    sheet=workbook['A-D.- FICHA']
    def values(address):
        return list(dict.fromkeys(str(c.value).strip() for row in sheet[address] for c in row if c.value is not None and str(c.value).strip()))
    catalogs={name:{'options':values(address),'source':f'A-D.- FICHA!{address}'} for name,address in LIST_RANGES.items()}
    # "Si tiene" and "NO HAY" are observed entries, not invented dropdown codes.
    catalogs['loadSignage']={'options':values('E326:E326'),'source':'A-D.- FICHA!E326 (observed value; full list unresolved)','allowOtherText':True}
    dump(ASSETS/'scap_options.json',{'sourceSha256':fingerprint,'catalogs':catalogs,
        'typesByCategory':{k:{'options':values(v),'source':f'A-D.- FICHA!{v}'} for k,v in TYPE_RANGES.items()}})
    definitions=json.loads((ROOT/'tools/scap_field_map.json').read_text(encoding='utf-8'))
    for field in definitions:
        field['label']=field.get('label') or str(sheet[field['labelCell']].value).strip().rstrip(':')
        field['source']='A-D.- FICHA!'+field['valueCell']
        if field['kind']=='CATALOG': assert field['catalog'] in catalogs,field
    dump(ASSETS/'scap_fields.json',{'sourceSha256':fingerprint,'fields':definitions})
    g=workbook['G.- CONDICION ESTADISTICA']; by_code={e['code']:e for e in elements}
    example=[]
    for row in range(14,28):
        code=str(int(g.cell(row,1).value));factor=g.cell(row,5).value
        assert factor==by_code[code]['importanceFactor'],(code,factor)
        example.append({'code':code,'quantity':g.cell(row,3).value,'importanceFactor':factor,
            'percentages':[g.cell(row,col).value or 0 for col in range(6,12)],
            'expectedCondition':g.cell(row+143,4).value,'source':f'G.- CONDICION ESTADISTICA!A{row}:K{row}'})
    dump(ROOT/'app/src/test/resources/scap_agua_blanca.json',{'sourceSha256':fingerprint,
        'elements':example,'expected':g['E191'].value,'classification':g['M12'].value.strip()})
    dump(ROOT/'docs/reference/SCAP_EXTRACTION.json',{'sourceSha256':fingerprint,'elementCount':len(elements),
        'brokenNames':{k:v.attr_text for k,v in formulas.defined_names.items() if '#REF!' in v.attr_text},
        'validations':[{'range':str(v.sqref),'formula':v.formula1} for v in formulas['A-D.- FICHA'].data_validations.dataValidation],
        'notes':['Whitespace trimmed; identical duplicate labels deduplicated, codes and factors unchanged.',
                 'Lists recovered from visible source tables when defined names are #REF!.',
                 'Category labels are PROVISIONALES and ARTESANALES, exactly as the workbook.',
                 'Reference observations are not defaults for new bridges.']})
    print(f'Extracted {len(elements)} unique elements and {len(catalogs)} option catalogs. SHA256 {fingerprint}')

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--source',type=Path,default=SOURCE)
    extract(parser.parse_args().source)
